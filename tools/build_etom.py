"""Reads: references/eTOM/GB921_Business_Process_Framework_Processes_Excel_v26.0.xlsx
(sheets `eTOM26.0` and `eTOM Deleted`) and references/eTOM/etom_v26.0.json
(secondary, cross-check only).
Writes: knowledge/etom/GB921.md (frontmatter only -- body preserved),
knowledge/etom/processes.json, knowledge/etom/deleted.json,
knowledge/index/etom-anomalies.md.
Track: assisted -- eTOM ships with TM Forum's Frameworx release train
(~2x/year, member-gated), not IG1228's 8-week cadence. Run after a new
workbook lands in references/eTOM/, then build_index.py (spec/spec-etom.md 7).

tools/build_etom.py -- turn the raw v26.0 eTOM Excel export into the
agent-facing corpus under knowledge/etom/ (spec/spec-etom.md 4).

The workbook is the source of record, NOT etom_v26.0.json: the JSON is a lossy
projection that drops every description, the UID, and the vertical group
(spec-etom.md 2). The JSON is loaded only to cross-check id/name/level/domain
after the build -- mismatches become an anomalies.md line, never a silent
overwrite in either direction.

Every build rule below exists because the raw data violates it somewhere
(spec-etom.md 6). The short version:
  - de-dupe on process id (the sheet lists a process once per Vertical Group);
    vertical_groups becomes the sorted union across those rows
  - decode HTML entities in every text field (&amp; &nbsp; &#183; &#233;)
  - backfill the 538 blank Domain cells: nearest ancestor row that has one,
    else the majority domain for the id's `1.X` prefix; flag domain_inferred
  - parent/children computed from the dotted id, not read from a column
    (id `1.1.1` is Level 2 -- level == parts - 1, so parent exists only at
    parts >= 4)
  - normalize placeholder / empty Extended Description ("Not used for this
    process element", "") to null
  - split a trailing "...was renamed... old name was X" note out of
    extended_description into renamed_from
  - normalize the Issue release tag to a bare version for added_in; the four
    rows carrying a description paragraph in the Issue column instead get
    added_in null and an anomalies.md line

Genuine upstream data-quality problems (duplicated ids, conflicting UIDs,
description text in the wrong column, JSON cross-check mismatches) are all
written to knowledge/index/etom-anomalies.md -- the same role
matrix-discrepancies.md plays for the use-case/component matrix: a tracked
fact reviewed once per refresh, not a lost observation.

Must be idempotent (spec/spec.md principle 7, success criterion in 9): run
twice on the same workbook, byte-identical output. Every list is sorted by a
hierarchical id key before being emitted; anomalies are sorted within each
section.

Usage:
    python build_etom.py [--origin URL] [--retrieved YYYY-MM-DD]
"""
import argparse
import datetime as _dt
import hashlib
import html
import json
import os
import re

import openpyxl

import _yaml_lite as yaml_lite

REPO_ROOT = os.path.join(os.path.dirname(__file__), "..")
REF_DIR = os.path.join(REPO_ROOT, "references", "eTOM")
XLSX_PATH = os.path.join(REF_DIR, "GB921_Business_Process_Framework_Processes_Excel_v26.0.xlsx")
JSON_PATH = os.path.join(REF_DIR, "etom_v26.0.json")

ETOM_DIR = os.path.join(REPO_ROOT, "knowledge", "etom")
PROCESSES_OUT = os.path.join(ETOM_DIR, "processes.json")
DELETED_OUT = os.path.join(ETOM_DIR, "deleted.json")
GB921_OUT = os.path.join(ETOM_DIR, "GB921.md")
ANOMALIES_OUT = os.path.join(REPO_ROOT, "knowledge", "index", "etom-anomalies.md")

# source.origin is a member-gated Frameworx download with no stable public URL;
# override with --origin once the real one is known (spec/tasks-etom.md 1.1).
DEFAULT_ORIGIN = "TODO: TM Forum Frameworx 26.0 -- GB921 Business Process Framework (eTOM) Excel export (member-gated)"
DEFAULT_RETRIEVED = "2026-09-08"  # date the workbook entered references/eTOM/ (commit 5ca41a1)

UNUSED_MARKER = "(This Process ID has been deliberately unused!)"
PLACEHOLDER_RE = re.compile(r"not used for this (process )?element\.?\s*$", re.I)


def is_placeholder(s):
    """The workbook's stock 'no real content here' filler, in its several
    punctuation variants ('Not used for this process element', with/without a
    trailing period, occasionally '... element')."""
    return s is not None and PLACEHOLDER_RE.fullmatch(s.strip()) is not None

# GB921.md body is a hand-maintained orientation doc (spec-etom.md 4.1, 9 Q1):
# the tool owns the frontmatter and seeds this body on first creation, then
# preserves whatever body is on disk on every subsequent run. Editing the prose
# below (or the file) is the intended way to keep it current -- it is the one
# deliberate exception to "don't hand-edit generated files" in this folder,
# the same way each other knowledge/ subfolder names its own static files.
DEFAULT_BODY = """\
# GB921 — Business Process Framework (eTOM), v26.0

The **Business Process Framework (eTOM)** is TM Forum's enterprise process
model for a service provider: a hierarchical catalogue of ~2,900 business
process elements (a.k.a. business activities), from broad Level-2 groupings
down to Level-7 leaf activities, each with a stable dotted identifier and a
`UID`.

This folder is the agent-facing view of the v26.0 export. It is **generated
from the Excel workbook** in `references/eTOM/` by `tools/build_etom.py` —
don't hand-edit the JSON files; re-run the tool. This `.md`'s prose body is
the exception: it is hand-maintained orientation, not generated data.

## What's here

| File | What it is |
|---|---|
| `processes.json` | the corpus — one entry per **distinct** process element (deduped across vertical groups), with `id`, `uid`, `name`, `level`, `domain` (`domain_inferred` when back-filled), `vertical_groups`, `parent`, `children`, `brief_description`, `extended_description`, `added_in`, `renamed_from` |
| `deleted.json` | the workbook's `eTOM Deleted` sheet — process ids removed at some point in eTOM's history (cumulative, **not** v26.0-only), so a stale id resolves to "was removed" rather than "unknown" |
| `../index/etom-anomalies.md` | data-quality findings in the raw export, reviewed once per refresh |
| `../index/etom-index.json` | `by_domain`, plus the eTOM↔ODA-component join (`implemented_by`) and the refs that don't resolve (`stale_component_refs`) — built by `tools/build_index.py` |

## Reading it

- **Identifier vs level.** The dotted id carries a leading framework root, so
  `1.1.1` is a **Level 2** process and `1.1.1.1` is Level 3 — `level == (dots
  in id) `. `parent` is `null` at Level 2 (its parent is the domain root,
  e.g. `1.1` "Market Domain", which is not itself a row).
- **`uid`** is eTOM's own stable numeric key. It survives renumbering better
  than the dotted id does — but it is not unique in this export (see
  `etom-anomalies.md`).
- **Domains** (Level-1): Market, Sales, Product, Customer, Service, Resource,
  Business Partner, Enterprise. `domain_inferred: true` means the workbook
  left the cell blank and the tool filled it from the process hierarchy —
  true for every process added in release 25.5 / 26.0.
- **Absence is signal.** A process id not in `processes.json` is genuinely
  not in eTOM 26.0. Check `deleted.json` before calling it unknown — and note
  eTOM was heavily **renumbered for 25.5/26.0**, so a `v24.0` eTOM reference
  in an ODA component's `component.yaml` will frequently not resolve here at
  all. `etom-index.json`'s `stale_component_refs` is the honest account of
  that gap; don't rewrite the components' references to paper over it.

## Provenance

See the frontmatter above. Source of record is the `.xlsx`; `etom_v26.0.json`
is a secondary file used only to cross-check the extraction.
"""


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def natkey(pid):
    """Hierarchical sort key: '1.1.2' < '1.1.10' < '1.2'."""
    return tuple(int(p) for p in str(pid).split("."))


def clean(s):
    """Decode HTML entities, normalize whitespace artefacts, strip. '' -> None."""
    if s is None:
        return None
    s = html.unescape(str(s)).replace("\xa0", " ").replace("\r\n", "\n")
    s = "\n".join(line.rstrip() for line in s.split("\n")).strip()
    return s or None


RENAME_LINE_RE = re.compile(r"(was renamed|old name was)", re.I)
OLD_NAME_RE = re.compile(r"old name was[:\s]+(.+?)\s*$", re.I)


def split_rename_note(ext):
    """(cleaned_extended_description, renamed_from|None). Pops trailing lines
    that are a 'this process was renamed ... old name was X' note."""
    if ext is None:
        return None, None
    renamed_from = None
    m = re.search(r"old name was[:\s]+(.+?)(?:[.\n]|$)", ext, re.I)
    if m:
        renamed_from = clean(m.group(1))
    lines = ext.split("\n")
    while lines and (not lines[-1].strip() or RENAME_LINE_RE.search(lines[-1])):
        if not lines[-1].strip():
            lines.pop()
            continue
        if RENAME_LINE_RE.search(lines[-1]):
            lines.pop()
        else:
            break
    cleaned = "\n".join(lines).strip() or None
    return cleaned, renamed_from


VERSION_RE = re.compile(r"(\d\d\.\d)\b")


def parse_added_in(issue, pid, anomalies):
    if issue is None:
        return None
    issue = str(issue).strip()
    if len(issue) > 40:
        anomalies["issue_column_misuse"].append(
            f"`{pid}` — Issue column holds a description paragraph, not a release tag"
        )
        return None
    m = VERSION_RE.search(issue)
    if m:
        return m.group(1)
    anomalies["unrecognized_issue_tag"].append(f"`{pid}` — Issue = {issue!r}")
    return None


def load_process_rows():
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb["eTOM26.0"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] is None:
            continue
        rows.append(
            {
                "name": r[0],
                "id": str(r[1]).strip(),
                "level": r[2],
                "extended": r[3],
                "brief": r[4],
                "domain": (r[5].strip() if isinstance(r[5], str) else r[5]) or None,
                "vertical_group": (r[6].strip() if isinstance(r[6], str) else r[6]) or None,
                "uid": r[7],
                "issue": r[8],
            }
        )
    wb.close()
    return rows


def load_deleted_rows():
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb["eTOM Deleted"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    return rows


def load_json_index():
    with open(JSON_PATH, encoding="utf-8") as f:
        doc = json.load(f)
    by_id = {}
    for e in doc.get("entries", []):
        by_id.setdefault(str(e["id"]).strip(), []).append(e)
    return doc.get("version"), by_id


def prefix_domain_majority(rows):
    from collections import Counter, defaultdict

    counts = defaultdict(Counter)
    for row in rows:
        if row["domain"]:
            counts[".".join(row["id"].split(".")[:2])][row["domain"]] += 1
    return {k: c.most_common(1)[0][0] for k, c in counts.items()}


def build_processes(rows, json_by_id, anomalies):
    from collections import defaultdict

    groups = defaultdict(list)
    for row in rows:
        groups[row["id"]].append(row)

    known_domain = {}
    for pid, grp in groups.items():
        d = next((g["domain"] for g in grp if g["domain"]), None)
        if d:
            known_domain[pid] = d
    majority = prefix_domain_majority(rows)

    def resolve_domain(pid):
        if pid in known_domain:
            return known_domain[pid], False
        parts = pid.split(".")
        for cut in range(len(parts) - 1, 1, -1):
            anc = ".".join(parts[:cut])
            if anc in known_domain:
                return known_domain[anc], True
        return majority.get(".".join(parts[:2])), True

    processes = {}
    for pid, grp in sorted(groups.items(), key=lambda kv: natkey(kv[0])):
        uids = sorted({g["uid"] for g in grp})
        canonical = next((g for g in grp if g["domain"]), grp[0])
        if len(uids) > 1:
            anomalies["conflicting_uid"].append(
                f"`{pid}` — {len(grp)} rows with different UIDs {uids}; kept "
                f"{canonical['uid']} (the row carrying a Domain)"
            )
        vgs = sorted({g["vertical_group"] for g in grp if g["vertical_group"]})
        domain, inferred = resolve_domain(pid)

        parts = pid.split(".")
        parent = ".".join(parts[:-1]) if len(parts) >= 4 else None

        ext = clean(canonical["extended"])
        if is_placeholder(ext):
            ext = None
        ext, renamed_from = split_rename_note(ext)
        if is_placeholder(ext):  # a rename note was hiding behind the filler
            ext = None
        brief = clean(canonical["brief"])
        if is_placeholder(brief):
            brief = None

        processes[pid] = {
            "id": pid,
            "uid": canonical["uid"],
            "name": clean(canonical["name"]),
            "level": canonical["level"],
            "domain": domain,
            "domain_inferred": inferred,
            "vertical_groups": vgs,
            "parent": parent,
            "children": [],
            "brief_description": brief,
            "extended_description": ext,
            "added_in": parse_added_in(canonical["issue"], pid, anomalies),
            "renamed_from": renamed_from,
        }

    # children, from present rows only
    for pid, p in processes.items():
        if p["parent"] and p["parent"] in processes:
            processes[p["parent"]]["children"].append(pid)
        elif p["parent"] and p["parent"] not in processes:
            anomalies["orphaned_process"].append(
                f"`{pid}` (Level {p['level']}) — parent `{p['parent']}` is not a row in the sheet"
            )
    for p in processes.values():
        p["children"].sort(key=natkey)

    # uid collisions across distinct ids
    from collections import defaultdict as _dd

    uid_to_ids = _dd(set)
    for p in processes.values():
        uid_to_ids[p["uid"]].add(p["id"])
    for uid, ids in sorted(uid_to_ids.items()):
        if len(ids) > 1:
            anomalies["uid_collision"].append(
                f"UID {uid} is shared by distinct processes {sorted(ids, key=natkey)}"
            )

    cross_check(processes, json_by_id, anomalies)
    return [processes[k] for k in sorted(processes, key=natkey)]


def cross_check(processes, json_by_id, anomalies):
    for pid, p in processes.items():
        entries = json_by_id.get(pid)
        if not entries:
            anomalies["json_mismatch"].append(f"`{pid}` — present in the workbook, absent from etom_v26.0.json")
            continue
        e = entries[0]
        jname = clean(e.get("name"))
        if jname and p["name"] and jname != p["name"]:
            anomalies["json_mismatch"].append(
                f"`{pid}` — name differs: workbook {p['name']!r} vs json {jname!r}"
            )
        if e.get("level") is not None and e["level"] != p["level"]:
            anomalies["json_mismatch"].append(
                f"`{pid}` — level differs: workbook {p['level']} vs json {e['level']}"
            )
        jdom = e.get("domain")
        if jdom and not p["domain_inferred"] and jdom != p["domain"]:
            anomalies["json_mismatch"].append(
                f"`{pid}` — domain differs: workbook {p['domain']!r} vs json {jdom!r}"
            )
    workbook_ids = set(processes)
    for jid in json_by_id:
        if jid not in workbook_ids:
            anomalies["json_mismatch"].append(f"`{jid}` — in etom_v26.0.json but not in the workbook sheet")


def build_deleted(rows, anomalies):
    seen = {}
    out = []
    for r in rows:
        raw_id = r[6]
        if raw_id is None:
            continue
        pid = str(raw_id).replace(UNUSED_MARKER, "").strip()
        if not re.fullmatch(r"\d+(\.\d+)*", pid):
            continue  # "deliberately unused" placeholder rows, not real deletions
        if pid in seen:
            continue
        seen[pid] = True
        issue = clean(r[13])
        deleted_in = None
        if issue:
            m = re.search(r"(\d+\.\d+)", issue)
            deleted_in = m.group(1) if m else None
        out.append(
            {
                "id": pid,
                "uid": r[1],
                "name": clean(r[0]),
                "level": r[4],
                "domain": clean(r[2]),
                "framework_status": clean(r[10]),
                "deleted_in": deleted_in,
                "original_process_identifier": clean(r[5]),
                "brief_description": clean(r[12]),
            }
        )
    out.sort(key=lambda d: natkey(d["id"]))
    return out


def flag_deleted_still_live(deleted, processes, anomalies):
    live = {p["id"] for p in processes}
    for d in deleted:
        if d["id"] in live:
            anomalies["deleted_but_live"].append(
                f"`{d['id']}` — listed on the eTOM Deleted sheet but also present as a live v26.0 process"
            )


ANOMALY_SECTIONS = [
    ("duplicated_ids", "Duplicated process ids (collapsed on build)"),
    ("conflicting_uid", "Ids appearing with conflicting UID / Domain"),
    ("uid_collision", "One UID shared by distinct process ids"),
    ("orphaned_process", "Processes whose parent id is missing from the sheet"),
    ("issue_column_misuse", "Description text found in the Issue column"),
    ("unrecognized_issue_tag", "Unrecognized Issue release tags"),
    ("deleted_but_live", "On the Deleted sheet but also live in v26.0"),
    ("json_mismatch", "etom_v26.0.json cross-check mismatches"),
]


def write_anomalies(anomalies, stats):
    lines = [
        "# eTOM v26.0 — data-quality anomalies",
        "",
        "Generated by `tools/build_etom.py`. Same role as `matrix-discrepancies.md`:",
        "findings in TM Forum's raw eTOM export that the build works around but that",
        "are worth reporting upstream and reviewing once per refresh (spec/spec-etom.md 4.4).",
        "Re-generated on every build — don't hand-edit; add commentary in `spec/tasks-etom.md`.",
        "",
        "## Build summary",
        "",
        f"- Source rows in `eTOM26.0` sheet: **{stats['raw_rows']}**",
        f"- Distinct process elements after de-dup: **{stats['processes']}**",
        f"- Domains back-filled (`domain_inferred: true`): **{stats['inferred_domains']}**",
        f"- `renamed_from` extracted: **{stats['renamed']}**",
        f"- Rows on the `eTOM Deleted` sheet kept: **{stats['deleted']}**",
        "",
    ]
    for key, title in ANOMALY_SECTIONS:
        items = sorted(set(anomalies.get(key, [])))
        lines.append(f"## {title} ({len(items)})")
        lines.append("")
        if items:
            lines.extend(f"- {it}" for it in items)
        else:
            lines.append("_None._")
        lines.append("")
    with open(ANOMALIES_OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines).rstrip() + "\n")


def write_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, sort_keys=False, ensure_ascii=False)
        f.write("\n")


def write_gb921(version, origin, retrieved):
    envelope = {
        "id": "GB921",
        "type": "etom",
        "name": "Business Process Framework (eTOM)",
        "version": version,
        "status": "Frameworx 26.0 - TM Forum published",
        "source": {
            "origin": origin,
            "license": "TM Forum (Frameworx / GB921)",
            "retrieved": retrieved,
            "sha256": sha256(XLSX_PATH),
            "raw_path": "../references/eTOM/GB921_Business_Process_Framework_Processes_Excel_v26.0.xlsx",
            "secondary": {
                "path": "../references/eTOM/etom_v26.0.json",
                "sha256": sha256(JSON_PATH),
            },
        },
        "links": {"components": []},
    }
    body = DEFAULT_BODY
    seeded = True
    if os.path.exists(GB921_OUT):
        with open(GB921_OUT, encoding="utf-8") as f:
            _, body = yaml_lite.split(f.read())
        seeded = False
    yaml_lite.write(GB921_OUT, envelope, body)
    return seeded


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--origin", default=DEFAULT_ORIGIN)
    ap.add_argument("--retrieved", default=DEFAULT_RETRIEVED)
    args = ap.parse_args()

    from collections import defaultdict

    anomalies = defaultdict(list)

    version, json_by_id = load_json_index()
    version = version or "v26.0"

    raw_rows = load_process_rows()

    # duplicated-id inventory (informational; collapse happens in build_processes)
    from collections import Counter

    for pid, n in sorted(Counter(r["id"] for r in raw_rows).items(), key=lambda kv: natkey(kv[0])):
        if n > 1:
            anomalies["duplicated_ids"].append(f"`{pid}` — {n} rows")

    processes = build_processes(raw_rows, json_by_id, anomalies)
    deleted = build_deleted(load_deleted_rows(), anomalies)
    flag_deleted_still_live(deleted, processes, anomalies)

    os.makedirs(ETOM_DIR, exist_ok=True)
    write_json(PROCESSES_OUT, processes)
    write_json(DELETED_OUT, deleted)
    seeded = write_gb921(version, args.origin, args.retrieved)

    stats = {
        "raw_rows": len(raw_rows),
        "processes": len(processes),
        "inferred_domains": sum(1 for p in processes if p["domain_inferred"]),
        "renamed": sum(1 for p in processes if p["renamed_from"]),
        "deleted": len(deleted),
    }
    write_anomalies(anomalies, stats)

    print(f"processes.json: {len(processes)} process elements")
    print(f"deleted.json:   {len(deleted)} deleted-process records")
    print(f"GB921.md:       frontmatter written (body {'seeded' if seeded else 'preserved'})")
    total_anoms = sum(len(set(v)) for v in anomalies.values())
    print(f"etom-anomalies.md: {total_anoms} finding(s) across {len([k for k in anomalies if anomalies[k]])} categories")


if __name__ == "__main__":
    main()
